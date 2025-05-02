from automation_manager import AutomationManager
from windia_ids import WindiaWindows, LN_ids , Windia_Buttons
from pywinauto import keyboard, mouse, findwindows, Application
import time
from docx import Document


def read_VE_fromdocx_file(path : str):
    patients = {}
    document = Document(path)
    for index, table in enumerate(document.tables):
        #print("Table", index)
        for row in range(len(table.rows)):
            for col in range(len(table.columns)):
                #print(table.cell(row, col).text, end='\t')
                pass
            
            LN_type = table.cell(row, 2).text if len(table.rows) > 2 else ""
            #print("LN_type" + str(LN_type))
            if LN_type == "E" or LN_type == "V" or LN_type == "E + V" or LN_type == "E+V"  :
                name = table.cell(row, 0).text
                name = name.split(",", 1)
                name = name[0] + ", " + name[1].split(' ', 2)[1]
                print(name)
                #print("name: " + str(name))
                patients[name] = LN_type
        print(patients)
        return patients
        



def input_hours_for_bill(automationManager : AutomationManager, max_hours : int, interships : int):
    offset = 20
    x =  600
    y =  380
    
    #click Selbstzahler tab
    automationManager.click_inside_window(WindiaWindows.LEISTUNGS_NACHWEIS, 6/9 , 1/8)
    
    
    click_distance_from_top_left(x,y, automationManager)
    keyboard.send_keys(str(max_hours))
    y += offset 
    #Click inside Table
    click_distance_from_top_left(x,y,automationManager)
    y += offset 
    
    if interships >= 2:
        click_distance_from_top_left(x,y,automationManager)
        y += offset 
        keyboard.send_keys(str(max_hours))

        click_distance_from_top_left(x,y,automationManager)
        y += offset 

    
    #click_distance_from_top_left( 3/10, 37/100, W.autoManager)


    #click Rechnug = Done
    #automationManager.click_inside_window(WindiaWindows.LEISTUNGS_NACHWEIS,3/11 , 9/10)

def click_distance_from_top_left(x : float , y : float , automationManager : AutomationManager , double = False): 
    
    rec = automationManager._get_rec(WindiaWindows.LEISTUNGS_NACHWEIS)
    print(rec)
    x_multiplier = 1
    y_multiplier = 1
    if (rec.left > 800):
        x_multiplier = 0.95
        y_multiplier = 0.81
    #automationManager.click_inside_window(WindiaWindows.LEISTUNGS_NACHWEIS, (x * x_multiplier), (y* y_multiplier))
    w = abs(rec.top - rec.bottom) * y
    if double:
        mouse.double_click(coords= ( int((rec.left + x) * x_multiplier) , int((rec.top + y) * y_multiplier)))
    else:
        mouse.click(coords= ( int((rec.left + x) * x_multiplier) , int((rec.top + y) * y_multiplier)))

def clickclick(automanager):
    offset = 20
    x =  600
    y =  370
    click_distance_from_top_left(x,y,automanager)
    y += offset 
    click_distance_from_top_left(x,y,automanager)

def find_caregiver_row(automanager):
    offset = 20
    x =  600
    y =  270
    ln_win = automanager.get_and_wait_for_window(WindiaWindows.LEISTUNGS_NACHWEIS, 3)

    found = False
    while not found:
        click_distance_from_top_left(x,y,automanager,True)
        print(y)
        try:
            btn = ln_win.child_window(auto_id=str(Windia_Buttons.LN_CANCEL_BTN.value), control_type="Button").wrapper_object()  # Button changed
            
            if btn:
                btn.invoke()
                found = True
                return y
        except:
            y += offset



def print_leistungsnachweis(automanager : AutomationManager, type, cur_page):
    #Entlastung tab
    automanager.click_inside_window(WindiaWindows.LEISTUNGS_NACHWEIS, 8/9 , 1/8)
    offset = 20
    x =  600
    y =  370
    ln_win = automanager.get_and_wait_for_window(WindiaWindows.LEISTUNGS_NACHWEIS, 3)

    if cur_page != type:        
        rec = ln_win.child_window(auto_id=str(LN_ids.UNDER_VH_BUTTON_EDIT.value), control_type="Edit").wrapper_object().rectangle()
        mouse.click(coords= (int(rec.left)  , int(rec.top - 30 )))
        y -= offset
        time.sleep(5)

    y = find_caregiver_row(automanager)
    y += 2*offset
    
    
    click_distance_from_top_left(x,y,automanager)
    y += offset 
    click_distance_from_top_left(x,y,automanager)

    
    #click Druck LN button
    ln_win = automanager.get_and_wait_for_window(WindiaWindows.LEISTUNGS_NACHWEIS, 3)
    rec_btn = ln_win.child_window(auto_id=str(Windia_Buttons.LN_EINSTELLUNG_BTN.value), control_type="Button").wrapper_object().rectangle()
    mouse.click(coords= (int(rec_btn.left) + 400 , int(rec_btn.top)))

    #print once
    time.sleep(4)
    ln_print_win =  automanager.get_and_wait_for_window(WindiaWindows.AUSDRUCK_LN, 3)
    #ln_print_win.child_window(auto_id= str(Windia_Buttons.LN_PRINT_ONCE_BTN.value), control_type="Button").wrapper_object().click()
    
    #close print page
    automanager.close_window(ln_print_win)
    automanager._click_popup_window_away("7", "WinDIA-Meldung")

    #close LN
    ln_win = automanager.get_and_wait_for_window(WindiaWindows.LEISTUNGS_NACHWEIS, 3)
    automanager.close_window(ln_win)
    automanager._click_popup_window_away("7", "WinDIA-Meldung")
    time.sleep(0.2)
    automanager._click_popup_window_away("7", "Leistungserfassung WinDIA", ln_win)
    time.sleep(0.2)

    
